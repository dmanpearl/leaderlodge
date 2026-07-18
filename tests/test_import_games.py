from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from django.contrib.auth.models import User
from openpyxl import Workbook

from leaderlodge.games.importing.progressiverummy import parse_progressiverummy_workbook
from leaderlodge.games.importing.service import (
    import_games_file,
    import_mahjongg_file,
    import_pickleball_file,
    import_progressiverummy_file,
    import_scorekeeper_file,
)
from leaderlodge.games.models import Game
from leaderlodge.groups.models import GroupMembership, GroupPlayer
from leaderlodge.groups.services import (
    add_guest_player,
    create_group_with_owner,
    resolve_seed_player,
)
from tests.fixtures.test_seeds import (
    EXPECTED_IMPORT_SUMMARIES,
    committed_test_seed_paths,
)


@pytest.fixture
def user(db):
    return User.objects.create_user(
        username="david",
        password="test-pass-123",
        first_name="David",
        last_name="Pearl",
    )


@pytest.fixture
def group(user):
    return create_group_with_owner(user=user, name="Family")


def _write_scorekeeper(
    path: Path,
    rows: list[tuple],
    *,
    headers: tuple[str, ...] = (
        "Date",
        "Nancy",
        "David",
        "Where",
        "Nancy",
        "David",
        "Diff",
    ),
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ScoreKeeper"
    worksheet.append(headers)
    worksheet.append(("Totals", 0, 0, None, 0, 0, None))
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


@pytest.fixture
def sample_spreadsheet(tmp_path):
    path = tmp_path / "dominos_seed_test.xlsx"
    _write_scorekeeper(
        path,
        [
            (datetime(2026, 3, 8), 25, 30, "Sally"),
            (datetime(2026, 3, 9), 30, 25, None),
        ],
    )
    return path


@pytest.mark.django_db
def test_resolve_seed_player_matches_linked_first_name(group, user):
    player = resolve_seed_player(group=group, seed_name="David")
    assert player.linked_user_id == user.pk


@pytest.mark.django_db
def test_resolve_seed_player_matches_existing_guest(group):
    guest = add_guest_player(group=group, display_name="Nancy")
    player = resolve_seed_player(group=group, seed_name="Nancy")
    assert player.pk == guest.pk


@pytest.mark.django_db
def test_resolve_seed_player_matches_inactive_guest(group):
    guest = add_guest_player(group=group, display_name="Nancy")
    guest.is_active = False
    guest.save(update_fields=["is_active"])

    player = resolve_seed_player(group=group, seed_name="Nancy")
    assert player.pk == guest.pk


@pytest.mark.django_db
def test_resolve_seed_player_matches_inactive_linked_first_name(group, user):
    nancy_user = User.objects.create_user(
        username="nmanpearl",
        password="test-pass-123",
        first_name="Nancy",
        last_name="Pearl",
    )
    player = GroupPlayer.objects.create(
        group=group,
        display_name="Nancy P.",
        linked_user=nancy_user,
        is_active=False,
    )

    resolved = resolve_seed_player(group=group, seed_name="Nancy")
    assert resolved.pk == player.pk
    assert resolved.linked_user_id == nancy_user.pk


@pytest.mark.django_db
def test_resolve_seed_player_creates_player_for_group_member_first_name(group, user):
    nancy_user = User.objects.create_user(
        username="nmanpearl",
        password="test-pass-123",
        first_name="Nancy",
    )
    GroupMembership.objects.create(
        group=group,
        user=nancy_user,
        role=GroupMembership.Role.MEMBER,
    )

    player = resolve_seed_player(group=group, seed_name="Nancy")
    assert player.linked_user_id == nancy_user.pk
    assert player.display_name == "Nancy"


@pytest.mark.django_db
def test_resolve_seed_player_prefers_linked_user_over_guest_display_name(group, user):
    add_guest_player(group=group, display_name="Nancy")
    nancy_user = User.objects.create_user(
        username="nmanpearl",
        password="test-pass-123",
        first_name="Nancy",
    )
    linked = GroupPlayer.objects.create(
        group=group,
        display_name="Nancy P.",
        linked_user=nancy_user,
    )

    player = resolve_seed_player(group=group, seed_name="Nancy")
    assert player.pk == linked.pk


@pytest.mark.django_db
def test_import_scorekeeper_file_uses_linked_nancy_player(
    group, user, sample_spreadsheet
):
    nancy_user = User.objects.create_user(
        username="nmanpearl",
        password="test-pass-123",
        first_name="Nancy",
    )
    nancy_player = GroupPlayer.objects.create(
        group=group,
        display_name="Nancy",
        linked_user=nancy_user,
    )

    summary = import_scorekeeper_file(
        path=sample_spreadsheet,
        group=group,
        game_type_code="dominos",
        recorded_by=user,
    )

    assert summary.created == 2
    assert GroupPlayer.objects.filter(group=group, display_name="Nancy").count() == 1
    games = Game.objects.filter(group=group).prefetch_related(
        "participant_results__player"
    )
    for game in games:
        nancy_participant = game.participant_results.get(player__linked_user=nancy_user)
        assert nancy_participant.player_id == nancy_player.pk


@pytest.mark.django_db
def test_resolve_seed_player_creates_guest(group):
    player = resolve_seed_player(group=group, seed_name="Nancy")
    assert player.display_name == "Nancy"
    assert player.linked_user is None


@pytest.mark.django_db
def test_import_scorekeeper_file_uses_header_player_names(group, user, tmp_path):
    path = tmp_path / "custom_players.xlsx"
    _write_scorekeeper(
        path,
        [
            (datetime(2026, 3, 8), 25, 30, "Lodge"),
        ],
        headers=("Date", "Carter", "Elena", "Where"),
    )
    summary = import_scorekeeper_file(
        path=path,
        group=group,
        game_type_code="dominos",
        recorded_by=user,
    )
    assert summary.created == 1
    assert GroupPlayer.objects.filter(group=group, display_name="Carter").exists()
    assert GroupPlayer.objects.filter(group=group, display_name="Elena").exists()
    assert not GroupPlayer.objects.filter(group=group, display_name="Nancy").exists()


@pytest.mark.django_db
def test_import_scorekeeper_file_is_idempotent(group, user, sample_spreadsheet):
    import_scorekeeper_file(
        path=sample_spreadsheet,
        group=group,
        game_type_code="dominos",
        recorded_by=user,
    )
    summary = import_scorekeeper_file(
        path=sample_spreadsheet,
        group=group,
        game_type_code="dominos",
        recorded_by=user,
    )

    assert summary.created == 0
    assert summary.skipped == 2
    assert Game.objects.filter(group=group).count() == 2


@pytest.mark.django_db
def test_import_committed_test_seeds(group, user):
    for game_type, path in committed_test_seed_paths(Path(".")):
        if not path.is_file():
            pytest.skip(f"Test seed missing: {path.name}")

        summary = import_games_file(
            path=path,
            group=group,
            game_type_code=game_type,
            recorded_by=user,
        )
        expected_parsed, expected_created, expected_skipped = EXPECTED_IMPORT_SUMMARIES[
            game_type
        ]
        assert summary.parsed == expected_parsed
        assert summary.created == expected_created
        assert summary.skipped == expected_skipped


def _write_progressiverummy(path: Path, rows: list[tuple]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Raw Games"
    worksheet.append(
        (
            "Date",
            "Where",
            "Half",
            "aa",
            "Nancy",
            "David",
            "Pat",
            "Min",
            "Max",
            "Winner",
            "Loser",
            "Players",
        )
    )
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


@pytest.fixture
def progressiverummy_spreadsheet(tmp_path):
    path = tmp_path / "progressiverummy_seed_test.xlsx"
    _write_progressiverummy(
        path,
        [
            (
                datetime(2026, 3, 8),
                "Sally",
                "Yes",
                "skip",
                10,
                20,
                30,
                10,
                30,
                "Nancy",
                "Pat",
                3,
            ),
            (
                datetime(2026, 3, 9),
                "<?>",
                None,
                None,
                5,
                15,
                25,
                5,
                25,
                "Nancy",
                "Pat",
                3,
            ),
            ("<?>", "Kitchen", None, None, 8, 18, 28, 8, 28, "Nancy", "Pat", 3),
        ],
    )
    return path


def test_parse_progressiverummy_workbook(progressiverummy_spreadsheet):
    rows = parse_progressiverummy_workbook(progressiverummy_spreadsheet)
    assert len(rows) == 3

    assert rows[0].played_on.isoformat() == "2026-03-08"
    assert rows[0].location_name == "Sally"
    assert rows[0].is_half is True
    assert [ps.name for ps in rows[0].player_scores] == ["Nancy", "David", "Pat"]
    assert rows[0].metadata["Winner"] == "Nancy"

    assert rows[1].location_name == ""
    assert rows[1].is_half is False

    assert rows[2].played_on.isoformat() == "2026-03-07"
    assert rows[2].location_name == "Kitchen"


@pytest.mark.django_db
def test_import_progressiverummy_file(group, user, progressiverummy_spreadsheet):
    summary = import_progressiverummy_file(
        path=progressiverummy_spreadsheet,
        group=group,
        recorded_by=user,
    )

    assert summary.parsed == 3
    assert summary.created == 3
    assert summary.skipped == 0

    games = Game.objects.filter(
        group=group, game_type_code="progressiverummy"
    ).order_by("pk")
    assert games.count() == 3
    assert games[0].half_weight == 2
    assert games[1].half_weight == 1
    assert games[2].half_weight == 1
    assert GroupPlayer.objects.filter(group=group, display_name="Pat").exists()


def _write_mahjongg(path: Path, rows: list[tuple]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Raw Games"
    worksheet.append(
        (
            "Date",
            "Where",
            "Half",
            "aa",
            "Carter",
            "Elena",
            "Lily",
            "Owen",
        )
    )
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


@pytest.fixture
def mahjongg_spreadsheet(tmp_path):
    path = tmp_path / "mahjongg_seed_test.xlsx"
    _write_mahjongg(
        path,
        [
            (datetime(2026, 3, 8), "Lodge", None, None, 25, 0, 0, 0),
            (datetime(2026, 3, 9), "The Club", None, None, 0, 30, 10, 5),
            (datetime(2026, 3, 10), "Treehouse", None, None, 10, 20, None, None),
        ],
    )
    return path


@pytest.mark.django_db
def test_import_mahjongg_file(group, user, mahjongg_spreadsheet):
    summary = import_mahjongg_file(
        path=mahjongg_spreadsheet,
        group=group,
        recorded_by=user,
    )

    assert summary.parsed == 3
    assert summary.created == 2
    assert summary.skipped == 1

    games = Game.objects.filter(group=group, game_type_code="mahjongg").order_by("pk")
    assert games.count() == 2
    assert all(game.half_weight == 1 for game in games)
    assert GroupPlayer.objects.filter(group=group, display_name="Owen").exists()

    winners = {
        game.participant_results.get(result="win").player.display_name for game in games
    }
    assert winners == {"Carter", "Elena"}


def _write_pickleball(path: Path, rows: list[tuple]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Raw Games"
    worksheet.append(
        (
            "Date",
            "Where",
            "Score type",
            "aa",
            "Harper+Owen",
            "Lucas",
            "Elena+Lucy",
        )
    )
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


@pytest.fixture
def pickleball_spreadsheet(tmp_path):
    path = tmp_path / "pickleball_seed_test.xlsx"
    _write_pickleball(
        path,
        [
            (datetime(2026, 7, 1), "Park A", None, None, 15, 10, None),
            (datetime(2026, 7, 2), "Park B", "Rally", None, 11, None, 0),
            (datetime(2026, 7, 3), "Park C", None, None, None, 17, 12),
            (datetime(2026, 7, 4), "Park D", None, None, 15, None, None),
        ],
    )
    return path


@pytest.mark.django_db
def test_import_pickleball_file(group, user, pickleball_spreadsheet):
    summary = import_pickleball_file(
        path=pickleball_spreadsheet,
        group=group,
        recorded_by=user,
    )

    assert summary.parsed == 4
    assert summary.created == 3
    assert summary.skipped == 1

    games = Game.objects.filter(group=group, game_type_code="pickleball").order_by("pk")
    assert games.count() == 3
    assert games.filter(score_type="rally").count() == 1
    assert GroupPlayer.objects.filter(group=group, is_team=True).exists()

    team_game = games.first()
    team_result = team_game.participant_results.filter(player__is_team=True).first()
    assert team_result is not None
    assert "+" in team_result.player.display_name


@pytest.mark.django_db
def test_import_committed_pickleball_demo_seed(group, user):
    path = Path("data/seeds/pickleball/pickleball_demo_2026-07-16.xlsx")
    if not path.is_file():
        pytest.skip("Pickleball demo seed not present")

    summary = import_pickleball_file(path=path, group=group, recorded_by=user)
    assert summary.parsed == 18
    assert summary.created == 15
    assert summary.skipped == 3
    assert Game.objects.filter(group=group, game_type_code="pickleball").count() == 15

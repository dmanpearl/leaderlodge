#!/usr/bin/env python3
"""Rewrite seed spreadsheet Where values with location cleanup rules."""

from __future__ import annotations

import re
import shutil
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SEED_ROOT = ROOT / "data" / "seeds"
OUT_STAMP = "2026-07-16"
CAFE_CUTOFF = date(2016, 6, 29)

# Source files to rewrite (old stamp -> new).
SOURCES = [
    "backgammon/backgammon_demo_2026-06-22.xlsx",
    "backgammon/backgammon_seed_2026-06-16.xlsx",
    "dominos/dominos_demo_2026-06-22.xlsx",
    "dominos/dominos_seed_2026-06-16.xlsx",
    "mahjongg/mahjongg_demo_2026-06-22.xlsx",
    "pickleball/pickleball_demo_2026-07-04.xlsx",
    "progressiverummy/progressiverummy_demo_2026-06-22.xlsx",
    "progressiverummy/progressiverummy_seed_2026-06-16.xlsx",
]


def fold_accents(name: str) -> str:
    decomposed = unicodedata.normalize("NFD", name)
    return "".join(ch for ch in decomposed if unicodedata.category(ch) != "Mn")


def norm(name: str) -> str:
    s = fold_accents(name.strip().lower())
    s = s.replace("'", "'").replace("'", "'").replace("`", "'")
    s = re.sub(r"\s+", " ", s)
    return s


def fuzzy(name: str) -> str:
    s = re.sub(r"[^a-z0-9\s']+", " ", norm(name))
    return re.sub(r"\s+", " ", s).strip()


def parse_played_on(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text or text in {"<?>", "Totals", "totals"}:
            return None
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def strip_site_numbers(name: str) -> str:
    """Drop campground/beach site numbers and trailing site labels."""
    s = name.strip()
    s = re.sub(r"\s*#\d+(\s*[-:][^\s,]*)?", "", s)
    s = re.sub(r"\s+\d+:\d+\b", "", s)  # e.g. Calico ... 9:9
    s = re.sub(r"\s*,\s*,", ",", s)
    s = re.sub(r"\s+", " ", s).strip(" ,-")
    return s or name.strip()


def is_home_variant(name: str) -> bool:
    f = fuzzy(name)
    if f == "home" or f.startswith("home "):
        return True
    # e.g. "Dining room table at home"
    return bool(re.search(r"\bat home\b", f))


def is_cafe_50(name: str) -> bool:
    f = fuzzy(name)
    # cafe 50's / cafe 50s / cafe 50 s
    return bool(re.search(r"\bcafe\s*50'?s?\b", f))


def is_ihop(name: str) -> bool:
    return "ihop" in fuzzy(name).replace("-", "")


def ihop_city(name: str) -> str | None:
    """Return explicit non-MDR city if present."""
    f = fuzzy(name)
    # after ihop, look for city tokens
    m = re.search(r"ihop\s*,?\s*(.+)$", f)
    if not m:
        return None
    rest = m.group(1).strip(" ,")
    if not rest:
        return None
    if rest in {"mdr", "marina del rey", "marina del rey ca"}:
        return None
    # Title-case-ish from original after comma if present
    if "," in name:
        return name.split(",", 1)[1].strip()
    return rest.title()


def is_j_nichols(name: str) -> bool:
    f = fuzzy(name)
    if re.search(r"\bj\.?\s*nichols?\b", f):
        return True
    return bool(re.fullmatch(r"nichols?(?:\s+mdr|\s+marina del rey)?", f))


def is_preston(name: str) -> bool:
    f = fuzzy(name)
    return f == "preston" or f.startswith("preston ")


def is_george(name: str) -> bool:
    f = fuzzy(name)
    return f.startswith("george")  # George's, Georges, etc.


def transform_where(name: str, played_on: date | None) -> str:
    original = name.strip()
    if not original:
        return original

    if is_home_variant(original):
        return "Preston Way"
    if is_preston(original):
        return "Preston Way"
    if is_george(original):
        return "George's Burgers, Santa Monica"
    if is_cafe_50(original):
        if played_on is not None and played_on >= CAFE_CUTOFF:
            return "Café 50's, WLA"
        return "Café 50's, Venice"
    if is_ihop(original):
        city = ihop_city(original)
        if city:
            return f"iHop, {city}"
        return "iHop, MDR"
    if is_j_nichols(original):
        return "J Nichols, MDR"

    return strip_site_numbers(original)


def find_where_col(header_row) -> int | None:
    for i, cell in enumerate(header_row):
        if cell is None:
            continue
        h = str(cell).strip().lower()
        if h in {"where", "location"}:
            return i
    return None


def out_path_for(src: Path) -> Path:
    name = src.name
    # Replace trailing _YYYY-MM-DD.xlsx with new stamp
    new_name = re.sub(r"_\d{4}-\d{2}-\d{2}\.xlsx$", f"_{OUT_STAMP}.xlsx", name)
    if new_name == name:
        new_name = src.stem + f"_{OUT_STAMP}.xlsx"
    return src.with_name(new_name)


def rewrite_file(src: Path) -> tuple[Path, int, Counter]:
    from collections import Counter

    dest = out_path_for(src)
    # Copy then edit so other sheets/styles preserved
    if dest.resolve() == src.resolve():
        raise SystemExit(f"Refusing to overwrite source in place: {src}")
    shutil.copy2(src, dest)

    wb = load_workbook(dest)
    changed = 0
    mapping: Counter[str] = Counter()
    sheet = None
    for candidate in ("ScoreKeeper", "Raw Games"):
        if candidate in wb.sheetnames:
            sheet = wb[candidate]
            break
    if sheet is None:
        sheet = wb[wb.sheetnames[0]]

    rows = sheet.iter_rows(min_row=1, max_row=1, values_only=True)
    header = next(rows, None)
    where_idx = find_where_col(header) if header else None
    if where_idx is None:
        wb.close()
        dest.unlink(missing_ok=True)
        raise SystemExit(f"No Where column in {src}")

    for row_cells in sheet.iter_rows(min_row=2):
        date_cell = row_cells[0].value
        if isinstance(date_cell, str) and date_cell.strip().lower() == "totals":
            continue
        where_cell = row_cells[where_idx]
        raw = where_cell.value
        if raw is None:
            continue
        text = str(raw).strip()
        if not text:
            continue
        played = parse_played_on(date_cell)
        new_val = transform_where(text, played)
        if new_val != text:
            where_cell.value = new_val
            changed += 1
            mapping[f"{text} => {new_val}"] += 1

    wb.save(dest)
    wb.close()
    return dest, changed, mapping


def main() -> None:
    from collections import Counter

    total_changed = 0
    all_map: Counter[str] = Counter()
    written: list[Path] = []
    for rel in SOURCES:
        src = SEED_ROOT / rel
        if not src.is_file():
            print(f"SKIP missing {src}")
            continue
        dest, changed, mapping = rewrite_file(src)
        written.append(dest)
        total_changed += changed
        all_map.update(mapping)
        print(f"{src.name} -> {dest.name}: {changed} Where cells updated")

    print(f"\nTotal cells updated: {total_changed}")
    print("\nTop replacements:")
    for key, cnt in all_map.most_common(40):
        print(f"  {cnt:4d}  {key}")

    print("\nWrote:")
    for p in written:
        print(f"  {p.relative_to(ROOT)}")


if __name__ == "__main__":
    main()

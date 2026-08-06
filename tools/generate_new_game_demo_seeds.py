#!/usr/bin/env python3
"""Generate demo seed spreadsheets with per-sport randomized size/span.

Ranges (randomized independently per game type):
  - game count: 50-80
  - history span: 3-6 years (ending mid-2026)

Run from repo root:
  source venv/bin/activate
  python tools/generate_new_game_demo_seeds.py
"""

from __future__ import annotations

import random
import sys
from datetime import date, timedelta
from itertools import combinations
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from leaderlodge.groups.demo_seed_venues_data import DEMO_SEED_VENUES  # noqa: E402

SEEDS = ROOT / "data" / "seeds"
SEED_DATE = "2026-08-06"
END = date(2026, 7, 28)
GAME_COUNT_RANGE = (50, 80)
YEAR_SPAN_RANGE = (3, 6)
LOCATIONS = tuple(row[0] for row in DEMO_SEED_VENUES)

PICKLEBALL_POOL = (
    "Amelia",
    "Owen",
    "Piper",
    "Wyatt",
    "Aurora",
    "Elena",
    "Lucy",
    "Hudson",
)


def _sport_params(rng: random.Random, label: str) -> tuple[int, date, list[date]]:
    """Pick game count and year span for one sport; return dates newest-bounded at END."""
    game_count = rng.randint(*GAME_COUNT_RANGE)
    years = rng.randint(*YEAR_SPAN_RANGE)
    start = date(END.year - years, END.month, END.day)
    if start.month == 2 and start.day == 29:
        start = date(start.year, 2, 28)
    span_days = (END - start).days
    days = sorted(rng.sample(range(span_days + 1), k=game_count))
    dates = [start + timedelta(days=d) for d in days]
    print(f"  {label}: {game_count} games over {years} years ({start} to {END})")
    return game_count, start, dates


def _write_raw_games(path: Path, players: list[str], rows: list[tuple]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Raw Games"
    worksheet.append(("Date", "Where", "Half", "aa", *players))
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    print(f"Wrote {path.relative_to(ROOT)} ({len(rows)} games)")


def _write_scorekeeper(
    path: Path, player_a: str, player_b: str, rows: list[tuple]
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ScoreKeeper"
    worksheet.append(("Date", player_a, player_b, "Where", player_a, player_b, "Diff"))
    worksheet.append(("Totals", None, None, None, None, None, None))
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    print(f"Wrote {path.relative_to(ROOT)} ({len(rows)} games)")


def _write_pickleball(path: Path, team_labels: list[str], rows: list[tuple]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Raw Games"
    worksheet.append(("Date", "Where", "Score type", "aa", *team_labels))
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    print(f"Wrote {path.relative_to(ROOT)} ({len(rows)} games)")


def _raw_row(
    played_on: date,
    location: str,
    players: list[str],
    scores: dict[str, int | None],
) -> tuple:
    return (
        played_on,
        location,
        None,
        None,
        *(scores.get(name) for name in players),
    )


def generate_ginrummy(rng: random.Random) -> int:
    players = ["Lucas", "Aurora", "Lily"]
    game_count, _start, dates = _sport_params(rng, "ginrummy")
    rows = []
    for played_on in dates:
        pair = rng.sample(players, 2)
        high = rng.randint(85, 220)
        low = rng.randint(20, high - 15)
        scores = (
            {pair[0]: high, pair[1]: low}
            if rng.random() < 0.5
            else {pair[0]: low, pair[1]: high}
        )
        rows.append(_raw_row(played_on, rng.choice(LOCATIONS), players, scores))
    _write_raw_games(
        SEEDS / "ginrummy" / f"ginrummy_demo_{SEED_DATE}.xlsx", players, rows
    )
    return game_count


def generate_cribbage(rng: random.Random) -> int:
    players = ["Lucas", "Lily", "Owen"]
    game_count, _start, dates = _sport_params(rng, "cribbage")
    n_shutout = max(1, round(game_count * 0.01))
    n_skunk = max(1, round(game_count * 0.03))
    n_double = max(1, round(game_count * 0.02))
    kinds = (
        ["shutout"] * n_shutout
        + ["skunk"] * n_skunk
        + ["double_skunk"] * n_double
        + ["normal"] * (game_count - n_shutout - n_skunk - n_double)
    )
    rng.shuffle(kinds)
    rows = []
    for played_on, kind in zip(dates, kinds, strict=True):
        n_players = 3 if rng.random() < 0.35 else 2
        table = rng.sample(players, n_players)
        winner = table[0]
        losers = table[1:]
        if kind == "shutout":
            scores = {winner: 121, **{name: 0 for name in losers}}
        elif kind == "double_skunk":
            scores = {
                winner: 121,
                **{name: rng.randint(20, 61) for name in losers},
            }
        elif kind == "skunk":
            scores = {
                winner: 121,
                **{name: rng.randint(62, 90) for name in losers},
            }
        else:
            scores = {
                winner: 121,
                **{name: rng.randint(95, 119) for name in losers},
            }
            if rng.random() < 0.08:
                scores[winner] = rng.randint(105, 120)
                for name in losers:
                    scores[name] = min(scores[name], scores[winner] - 5)
        rows.append(_raw_row(played_on, rng.choice(LOCATIONS), players, scores))
    _write_raw_games(
        SEEDS / "cribbage" / f"cribbage_demo_{SEED_DATE}.xlsx", players, rows
    )
    return game_count


def generate_phase10(rng: random.Random) -> int:
    players = ["Lucas", "Aurora", "Harper", "Elena"]
    game_count, _start, dates = _sport_params(rng, "phase10")
    rows = []
    for played_on in dates:
        n = rng.choice([3, 3, 4, 4, 4])
        table = rng.sample(players, n)
        winner_score = rng.randint(15, 75)
        scores: dict[str, int | None] = {table[0]: winner_score}
        for name in table[1:]:
            scores[name] = winner_score + rng.randint(20, 160)
        rows.append(_raw_row(played_on, rng.choice(LOCATIONS), players, scores))
    _write_raw_games(
        SEEDS / "phase10" / f"phase10_demo_{SEED_DATE}.xlsx", players, rows
    )
    return game_count


def generate_darts(rng: random.Random) -> int:
    players = ["Cooper", "Lucy", "Mason", "Harper"]
    game_count, _start, dates = _sport_params(rng, "darts")
    rows = []
    for played_on in dates:
        n = rng.choice([2, 2, 3])
        table = rng.sample(players, n)
        ordered = sorted(
            ((name, rng.randint(18, 95)) for name in table),
            key=lambda item: item[1],
            reverse=True,
        )
        if len(ordered) > 1 and ordered[0][1] == ordered[1][1]:
            ordered[0] = (ordered[0][0], ordered[0][1] + 3)
        rows.append(_raw_row(played_on, rng.choice(LOCATIONS), players, dict(ordered)))
    _write_raw_games(SEEDS / "darts" / f"darts_demo_{SEED_DATE}.xlsx", players, rows)
    return game_count


def generate_cornhole(rng: random.Random) -> int:
    players = ["Harper", "Owen", "Elena"]
    game_count, _start, dates = _sport_params(rng, "cornhole")
    rows = []
    for played_on in dates:
        pair = rng.sample(players, 2)
        winner_score = 21
        loser_score = rng.randint(3, 19)
        scores = (
            {pair[0]: winner_score, pair[1]: loser_score}
            if rng.random() < 0.5
            else {pair[0]: loser_score, pair[1]: winner_score}
        )
        rows.append(_raw_row(played_on, rng.choice(LOCATIONS), players, scores))
    _write_raw_games(
        SEEDS / "cornhole" / f"cornhole_demo_{SEED_DATE}.xlsx", players, rows
    )
    return game_count


def generate_backgammon(rng: random.Random) -> int:
    players = ["Carter", "Elena"]
    game_count, _start, dates = _sport_params(rng, "backgammon")
    rows = []
    for played_on in dates:
        # Match points: 1 normal, 2 gammon, rare 3 backgammon.
        points = rng.choices([1, 2, 3], weights=[75, 20, 5], k=1)[0]
        if rng.random() < 0.5:
            a, b = points, 0
        else:
            a, b = 0, points
        rows.append((played_on, a, b, rng.choice(LOCATIONS)))
    _write_scorekeeper(
        SEEDS / "backgammon" / f"backgammon_demo_{SEED_DATE}.xlsx",
        players[0],
        players[1],
        rows,
    )
    return game_count


def generate_mahjongg(rng: random.Random) -> int:
    players = ["Carter", "Elena", "Lily", "Owen", "Mason"]
    game_count, _start, dates = _sport_params(rng, "mahjongg")
    rows = []
    for played_on in dates:
        n = rng.choice([3, 4, 4, 4])
        table = rng.sample(players, n)
        winner = table[0]
        # High wins; wall-ish zeros for others are common.
        scores: dict[str, int | None] = {winner: rng.choice([25, 30, 35, 45, 50])}
        for name in table[1:]:
            scores[name] = rng.choice([0, 0, 0, 5, 10])
        rows.append(_raw_row(played_on, rng.choice(LOCATIONS), players, scores))
    _write_raw_games(
        SEEDS / "mahjongg" / f"mahjongg_demo_{SEED_DATE}.xlsx", players, rows
    )
    return game_count


def generate_pickleball(rng: random.Random) -> int:
    # 40% of the 8-name pool is ~3; doubles needs 4, so use 4 (50%).
    n_roster = max(4, round(len(PICKLEBALL_POOL) * 0.40))
    roster = sorted(rng.sample(list(PICKLEBALL_POOL), n_roster))
    print(
        f"  pickleball roster ({n_roster}/{len(PICKLEBALL_POOL)}): {', '.join(roster)}"
    )

    team_labels = [f"{a}+{b}" for a, b in combinations(roster, 2)]
    game_count, _start, dates = _sport_params(rng, "pickleball")
    n_shutout = max(1, round(game_count * 0.05))
    shutout_flags = [True] * n_shutout + [False] * (game_count - n_shutout)
    rng.shuffle(shutout_flags)

    rows = []
    for played_on, is_shutout in zip(dates, shutout_flags, strict=True):
        # Two disjoint pairs from the 4-player roster.
        shuffled = roster[:]
        rng.shuffle(shuffled)
        team_a = "+".join(sorted(shuffled[:2]))
        team_b = "+".join(sorted(shuffled[2:4]))
        # Canonical label order matches header construction (combinations order).
        # Rebuild with sorted member names already; ensure header spelling.
        if team_a not in team_labels:
            team_a = "+".join(sorted(team_a.split("+")))
        if team_b not in team_labels:
            team_b = "+".join(sorted(team_b.split("+")))

        score_type = "Rally" if rng.random() < 0.35 else None
        target = 15 if score_type == "Rally" else 11
        if is_shutout:
            winner_score = target
            loser_score = 0
        else:
            winner_score = target
            loser_score = rng.randint(2, target - 2)

        if rng.random() < 0.5:
            scores = {team_a: winner_score, team_b: loser_score}
        else:
            scores = {team_a: loser_score, team_b: winner_score}

        rows.append(
            (
                played_on,
                rng.choice(LOCATIONS),
                score_type,
                None,
                *(scores.get(label) for label in team_labels),
            )
        )

    _write_pickleball(
        SEEDS / "pickleball" / f"pickleball_demo_{SEED_DATE}.xlsx",
        team_labels,
        rows,
    )
    return game_count


def main() -> None:
    rng = random.Random(20260806)
    print("Generating demo seeds (per-sport randomized counts/spans):")
    generate_ginrummy(rng)
    generate_cribbage(rng)
    generate_phase10(rng)
    generate_darts(rng)
    generate_cornhole(rng)
    generate_backgammon(rng)
    generate_mahjongg(rng)
    generate_pickleball(rng)


if __name__ == "__main__":
    main()

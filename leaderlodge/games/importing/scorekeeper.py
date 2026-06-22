from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

SCOREKEEPER_SHEET = "ScoreKeeper"
HEADER_ROW = 1
DATA_START_ROW = 3
DATE_COL = 0
PLAYER_ONE_COL = 1
PLAYER_TWO_COL = 2
WHERE_COL = 3


@dataclass(frozen=True)
class ScoreKeeperRow:
    row_number: int
    played_on: date
    scores: tuple[int, int]
    location_name: str


@dataclass(frozen=True)
class ScoreKeeperParseResult:
    player_names: tuple[str, str]
    rows: list[ScoreKeeperRow]


def _normalize_header(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _coerce_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _coerce_score(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _coerce_location(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_scorekeeper_workbook(path: Path) -> ScoreKeeperParseResult:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SCOREKEEPER_SHEET not in workbook.sheetnames:
            raise ValueError(
                f'Expected a "{SCOREKEEPER_SHEET}" worksheet in {path.name}.'
            )
        worksheet = workbook[SCOREKEEPER_SHEET]
        header_row = next(
            worksheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True)
        )
        player_one_name = _normalize_header(
            header_row[PLAYER_ONE_COL] if len(header_row) > PLAYER_ONE_COL else None
        )
        player_two_name = _normalize_header(
            header_row[PLAYER_TWO_COL] if len(header_row) > PLAYER_TWO_COL else None
        )
        if not player_one_name or not player_two_name:
            raise ValueError(
                f"Missing player name headers in columns B and C of {path.name}."
            )

        rows: list[ScoreKeeperRow] = []
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=DATA_START_ROW, values_only=True),
            start=DATA_START_ROW,
        ):
            if not values:
                continue

            played_on = _coerce_date(values[DATE_COL])
            if played_on is None:
                continue

            player_one_score = _coerce_score(values[PLAYER_ONE_COL])
            player_two_score = _coerce_score(values[PLAYER_TWO_COL])
            if player_one_score is None or player_two_score is None:
                continue

            rows.append(
                ScoreKeeperRow(
                    row_number=row_number,
                    played_on=played_on,
                    scores=(player_one_score, player_two_score),
                    location_name=_coerce_location(
                        values[WHERE_COL] if len(values) > WHERE_COL else None
                    ),
                )
            )
        return ScoreKeeperParseResult(
            player_names=(player_one_name, player_two_name),
            rows=rows,
        )
    finally:
        workbook.close()

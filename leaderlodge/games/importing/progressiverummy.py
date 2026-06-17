from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

RAW_GAMES_SHEET = "Raw Games"
DATA_START_ROW = 2
UNKNOWN_VALUE = "<?>"
METADATA_COLUMNS = frozenset({"Min", "Max", "Winner", "Loser", "Players"})
PLACEHOLDER_COLUMNS = frozenset({"aa"})


@dataclass(frozen=True)
class PlayerScore:
    name: str
    score: int


@dataclass
class ProgressiveRummyRow:
    row_number: int
    played_on: date | None
    location_name: str
    is_half: bool
    player_scores: list[PlayerScore] = field(default_factory=list)
    metadata: dict[str, object] = field(default_factory=dict)


def _normalize_header(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_unknown(value) -> bool:
    if value is None:
        return True
    return str(value).strip() == UNKNOWN_VALUE


def _coerce_date(value) -> date | None:
    if _is_unknown(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _coerce_score(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _coerce_location(value) -> str:
    if _is_unknown(value):
        return ""
    return str(value).strip()


def _is_half_game(value) -> bool:
    if value is None:
        return False
    return str(value).strip().lower() == "yes"


def _parse_columns(header_row: tuple) -> dict[str, int | list[tuple[int, str]]]:
    columns: dict[str, int | list[tuple[int, str]]] = {
        "date": None,
        "location": None,
        "half": None,
        "players": [],
    }
    for index, raw_name in enumerate(header_row):
        name = _normalize_header(raw_name)
        if not name or name in PLACEHOLDER_COLUMNS or name in METADATA_COLUMNS:
            continue
        if name == "Date":
            columns["date"] = index
        elif name == "Where":
            columns["location"] = index
        elif name == "Half":
            columns["half"] = index
        else:
            columns["players"].append((index, name))
    return columns


def _assign_undated_played_on(rows: list[ProgressiveRummyRow]) -> None:
    dated = [row.played_on for row in rows if row.played_on is not None]
    if not dated:
        fallback = date(1900, 1, 1)
        for row in rows:
            if row.played_on is None:
                row.played_on = fallback
        return

    earliest = min(dated)
    anchor = earliest - timedelta(days=1)
    for row in rows:
        if row.played_on is None:
            row.played_on = anchor


def parse_progressiverummy_workbook(path: Path) -> list[ProgressiveRummyRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if RAW_GAMES_SHEET not in workbook.sheetnames:
            raise ValueError(
                f'Expected a "{RAW_GAMES_SHEET}" worksheet in {path.name}.'
            )
        worksheet = workbook[RAW_GAMES_SHEET]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        columns = _parse_columns(header_row)

        if columns["date"] is None:
            raise ValueError(f'Missing "Date" column in {path.name}.')
        if columns["location"] is None:
            raise ValueError(f'Missing "Where" column in {path.name}.')
        if columns["half"] is None:
            raise ValueError(f'Missing "Half" column in {path.name}.')
        if not columns["players"]:
            raise ValueError(f'No player columns found in {path.name}.')

        player_columns = columns["players"]
        metadata_indexes = {
            index
            for index, name in enumerate(header_row)
            if _normalize_header(name) in METADATA_COLUMNS
        }

        parsed_rows: list[ProgressiveRummyRow] = []
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=DATA_START_ROW, values_only=True),
            start=DATA_START_ROW,
        ):
            if not values:
                continue

            player_scores: list[PlayerScore] = []
            for col_index, player_name in player_columns:
                score = _coerce_score(values[col_index] if col_index < len(values) else None)
                if score is None:
                    continue
                player_scores.append(PlayerScore(name=player_name, score=score))

            if not player_scores:
                continue

            raw_date = values[columns["date"]] if columns["date"] < len(values) else None
            raw_location = (
                values[columns["location"]] if columns["location"] < len(values) else None
            )
            raw_half = values[columns["half"]] if columns["half"] < len(values) else None

            metadata = {}
            for col_index in metadata_indexes:
                if col_index < len(values):
                    header_name = _normalize_header(header_row[col_index])
                    metadata[header_name] = values[col_index]

            parsed_rows.append(
                ProgressiveRummyRow(
                    row_number=row_number,
                    played_on=_coerce_date(raw_date),
                    location_name=_coerce_location(raw_location),
                    is_half=_is_half_game(raw_half),
                    player_scores=player_scores,
                    metadata=metadata,
                )
            )

        _assign_undated_played_on(parsed_rows)
        return parsed_rows
    finally:
        workbook.close()

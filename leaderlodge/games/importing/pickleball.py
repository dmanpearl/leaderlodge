from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from leaderlodge.games.registry import ScoreType
from leaderlodge.groups.models import Group, GroupPlayer
from leaderlodge.groups.services import resolve_seed_player
from leaderlodge.groups.teams import resolve_side_player

RAW_GAMES_SHEET = "Raw Games"
DATA_START_ROW = 2
UNKNOWN_VALUE = "<?>"
PLACEHOLDER_COLUMNS = frozenset({"aa"})
METADATA_COLUMNS = frozenset(
    {
        "Min",
        "Max",
        "Winner",
        "Loser",
        "Players",
        "Wins",
        "Losses",
    }
)


@dataclass(frozen=True)
class SideScore:
    label: str
    score: int


@dataclass
class PickleballRow:
    row_number: int
    played_on: date
    location_name: str
    score_type: str
    side_scores: list[SideScore] = field(default_factory=list)


def _normalize_header(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_unknown(value) -> bool:
    if value is None:
        return True
    return str(value).strip() == UNKNOWN_VALUE


def _coerce_date(value) -> date | None:
    if _is_unknown(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _coerce_score(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _coerce_location(value) -> str:
    if _is_unknown(value):
        return ""
    return str(value).strip()


def _coerce_score_type(value) -> str:
    if value is None:
        return ScoreType.TRADITIONAL.value
    cleaned = str(value).strip()
    if not cleaned:
        return ScoreType.TRADITIONAL.value
    if cleaned.lower() == ScoreType.RALLY.value:
        return ScoreType.RALLY.value
    return ScoreType.TRADITIONAL.value


def _parse_columns(header_row: tuple) -> dict[str, int | list[tuple[int, str]]]:
    columns: dict[str, int | list[tuple[int, str]]] = {
        "date": None,
        "location": None,
        "score_type": None,
        "sides": [],
    }
    for index, raw_name in enumerate(header_row):
        name = _normalize_header(raw_name)
        if not name or name in PLACEHOLDER_COLUMNS or name in METADATA_COLUMNS:
            continue
        if name == "Date":
            columns["date"] = index
        elif name == "Where":
            columns["location"] = index
        elif name == "Score type":
            columns["score_type"] = index
        elif name == "Half":
            continue
        else:
            columns["sides"].append((index, name))
    return columns


def parse_pickleball_workbook(path: Path) -> list[PickleballRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if RAW_GAMES_SHEET not in workbook.sheetnames:
            raise ValueError(
                f'Expected a "{RAW_GAMES_SHEET}" worksheet in {path.name}.'
            )
        worksheet = workbook[RAW_GAMES_SHEET]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        columns = _parse_columns(header_row)

        if columns["date"] is None:
            raise ValueError(f'Missing "Date" column in {path.name}.')
        if columns["location"] is None:
            raise ValueError(f'Missing "Where" column in {path.name}.')
        if not columns["sides"]:
            raise ValueError(f"No side columns found in {path.name}.")

        side_columns = columns["sides"]
        parsed_rows: list[PickleballRow] = []
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=DATA_START_ROW, values_only=True),
            start=DATA_START_ROW,
        ):
            if not values:
                continue

            played_on = _coerce_date(
                values[columns["date"]] if columns["date"] < len(values) else None
            )
            if played_on is None:
                continue

            side_scores: list[SideScore] = []
            for col_index, side_label in side_columns:
                score = _coerce_score(
                    values[col_index] if col_index < len(values) else None
                )
                if score is None:
                    continue
                side_scores.append(SideScore(label=side_label, score=score))

            if not side_scores:
                continue

            raw_location = (
                values[columns["location"]]
                if columns["location"] < len(values)
                else None
            )
            raw_score_type = (
                values[columns["score_type"]]
                if columns["score_type"] is not None
                and columns["score_type"] < len(values)
                else None
            )

            parsed_rows.append(
                PickleballRow(
                    row_number=row_number,
                    played_on=played_on,
                    location_name=_coerce_location(raw_location),
                    score_type=_coerce_score_type(raw_score_type),
                    side_scores=side_scores,
                )
            )

        return parsed_rows
    finally:
        workbook.close()


def _side_member_seed_names(side_label: str) -> list[str]:
    cleaned = side_label.strip()
    if "+" in cleaned:
        return [part.strip() for part in cleaned.split("+") if part.strip()]
    return [cleaned]


def sides_have_shared_members(side_labels: list[str]) -> bool:
    seen: set[str] = set()
    for label in side_labels:
        for name in _side_member_seed_names(label):
            if name in seen:
                return True
            seen.add(name)
    return False


def resolve_import_side(
    *,
    group: Group,
    side_label: str,
    cache: dict[str, GroupPlayer],
) -> GroupPlayer:
    cleaned = side_label.strip()
    if not cleaned:
        raise ValueError("Side label cannot be empty.")
    if cleaned in cache:
        return cache[cleaned]

    if "+" in cleaned:
        member_names = [part.strip() for part in cleaned.split("+") if part.strip()]
        if len(member_names) < 2:
            raise ValueError(f"Invalid team label: {cleaned!r}.")
        members = [
            resolve_seed_player(group=group, seed_name=name) for name in member_names
        ]
        player = resolve_side_player(group=group, members=members)
    else:
        player = resolve_seed_player(group=group, seed_name=cleaned)

    cache[cleaned] = player
    return player
